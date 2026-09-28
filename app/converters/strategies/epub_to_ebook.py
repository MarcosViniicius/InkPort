"""EPUB -> other ebook formats, in pure Python (KEPUB, TXT, FB2, DOCX).

These replace the Calibre conversions for the non-Kindle, non-PDF targets.
Everything runs in-process: no executable, nothing to install.
"""

from __future__ import annotations

import io
import posixpath
import re
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

from lxml import etree
from lxml import html as lxml_html

from app.converters.base import BaseConverter, ConversionError, ConversionRequest, ConversionResult
from app.converters.epub.reader import EpubDocument, read_epub

#: Sentence enders used to split text into Kobo spans.
_SENTENCE_END = re.compile(r"(?<=[.!?…])\s+")
_UNWANTED = ("script", "style", "head", "title", "meta", "link")


def _output(request: ConversionRequest, name: str) -> Path:
    """Output path inside the job workdir (created if the caller did not)."""
    request.workdir.mkdir(parents=True, exist_ok=True)
    return request.workdir / name


class EpubToKepubConverter(BaseConverter):
    """EPUB -> KEPUB: the same book with Kobo's sentence spans."""

    name = "epub_to_kepub"
    priority = 62

    def can_handle(self, request: ConversionRequest) -> bool:
        return request.target_format == "kepub" and request.detection.format == "epub"

    def supports_pair(self, source_fmt: str, target_fmt: str) -> bool:
        return source_fmt == "epub" and target_fmt == "kepub"

    def run(self, request: ConversionRequest) -> ConversionResult:
        out = _output(request, "output.kepub.epub")
        request.report(10, "Preparando o KEPUB")
        try:
            with zipfile.ZipFile(request.source) as src, zipfile.ZipFile(out, "w") as dst:
                # The mimetype entry must come first and stay uncompressed, and
                # the Kobo reader is strict about it.
                if "mimetype" in src.namelist():
                    dst.writestr(
                        "mimetype", src.read("mimetype"), compress_type=zipfile.ZIP_STORED
                    )
                documents = [
                    n for n in src.namelist() if n.lower().endswith((".xhtml", ".html", ".htm"))
                ]
                for index, name in enumerate(documents, start=1):
                    raw = src.read(name)
                    dst.writestr(name, _rewrite_document(raw, index))
                for name in src.namelist():
                    if name in documents or name == "mimetype" or name.endswith("/"):
                        continue
                    dst.writestr(name, src.read(name), compress_type=zipfile.ZIP_DEFLATED)
        except (zipfile.BadZipFile, OSError) as exc:
            raise ConversionError(f"EPUB inválido para KEPUB: {exc}") from exc

        request.report(100, "KEPUB pronto")
        return ConversionResult(
            output_path=out,
            converter=self.name,
            message="Convertido para KEPUB (spans do Kobo)",
        )


#: Elements whose text becomes Kobo spans (Kobo counts progress over them).
_TEXT_BLOCKS = {
    "p", "div", "h1", "h2", "h3", "h4", "h5", "h6", "li", "td", "blockquote",
}


def _rewrite_document(raw: bytes, file_index: int) -> bytes:
    try:
        root = lxml_html.document_fromstring(raw)
    except etree.ParserError:
        return raw
    counter = _wrap_tree(root, file_index)
    if counter == 0:
        return raw
    return lxml_html.tostring(root, encoding="utf-8", method="xml")


def _wrap_tree(root, file_index: int) -> int:
    """One ``span.koboSpan`` per sentence, in document order."""
    counter = 0
    for element in list(root.iter()):
        if not isinstance(element.tag, str) or element.tag.lower() in _UNWANTED:
            continue
        if element.tag.lower() not in _TEXT_BLOCKS:
            continue
        counter = _wrap_text(element, element.text, None, file_index, counter)
        for child in list(element):
            if child.tail and child.tail.strip():
                counter = _wrap_text(element, child.tail, child, file_index, counter)
    return counter


def _wrap_text(parent, text: str | None, anchor, file_index: int, counter: int) -> int:
    """Replace one text run by sentence spans, after ``anchor`` (or inline)."""
    if not text or not text.strip():
        return counter
    sentences = [chunk for chunk in _SENTENCE_END.split(text) if chunk]
    if not sentences:
        return counter
    if anchor is None:
        parent.text = None
    else:
        anchor.tail = None
    for offset, sentence in enumerate(sentences):
        counter += 1
        span = etree.Element("span")
        span.set("class", "koboSpan")
        span.set("id", f"kobo.{file_index}.{counter}")
        span.text = sentence if sentence.endswith(" ") else sentence + " "
        if anchor is None and offset == 0:
            parent.insert(0, span)
        elif anchor is None:
            parent.insert(offset, span)
        else:
            anchor.addnext(span)
            anchor = span
    return counter


class EpubToTextConverter(BaseConverter):
    """EPUB -> plain text, chapter per chapter."""

    name = "epub_to_text"
    priority = 60

    def can_handle(self, request: ConversionRequest) -> bool:
        return request.target_format == "txt" and request.detection.format == "epub"

    def supports_pair(self, source_fmt: str, target_fmt: str) -> bool:
        return source_fmt == "epub" and target_fmt == "txt"

    def run(self, request: ConversionRequest) -> ConversionResult:
        document = _load(request)
        chapters = epub_to_chapters(document)
        out = _output(request, "output.txt")
        parts: list[str] = []
        for title, body in chapters:
            if title:
                parts.append(f"\n\n{'=' * 4} {title} {'=' * 4}\n")
            parts.append(body)
        out.write_text("\n".join(parts).strip() + "\n", encoding="utf-8")
        request.report(100, "TXT pronto")
        return ConversionResult(output_path=out, converter=self.name, message="Convertido para TXT")


class EpubToFb2Converter(BaseConverter):
    """EPUB -> FB2 (an XML format FictionBook readers understand)."""

    name = "epub_to_fb2"
    priority = 60

    def can_handle(self, request: ConversionRequest) -> bool:
        return request.target_format == "fb2" and request.detection.format == "epub"

    def supports_pair(self, source_fmt: str, target_fmt: str) -> bool:
        return source_fmt == "epub" and target_fmt == "fb2"

    def run(self, request: ConversionRequest) -> ConversionResult:
        document = _load(request)
        chapters = epub_to_chapters(document, keep_markup=True)
        out = _output(request, "output.fb2")
        meta = request.metadata
        title = escape(meta.title or document.title or "Sem título")
        author_parts = (meta.author or document.author or "").split()
        first = escape(author_parts[0]) if author_parts else ""
        last = escape(" ".join(author_parts[1:])) if len(author_parts) > 1 else ""

        body: list[str] = []
        for chapter_title, xhtml in chapters:
            section = ["<section>"]
            if chapter_title:
                section.append(f"<title><p>{escape(chapter_title)}</p></title>")
            section.append(_xhtml_to_fb2(xhtml))
            section.append("</section>")
            body.append("".join(section))

        payload = (
            '<?xml version="1.0" encoding="utf-8"?>\n'
            '<FictionBook xmlns="http://www.gribuser.ru/xml/fictionbook/2.0"'
            ' xmlns:l="http://www.w3.org/1999/xlink">\n'
            "<description><title-info>"
            f"<genre>fiction</genre><author><first-name>{first}</first-name>"
            f"<last-name>{last}</last-name></author><book-title>{title}</book-title>"
            f"<lang>{escape(meta.language or document.language or 'pt')}</lang>"
            "</title-info>"
            f"<document-info><program-used>opds-server</program-used></document-info>"
            "</description>\n"
            f"<body>{''.join(body)}</body>\n"
            "</FictionBook>\n"
        )
        out.write_text(payload, encoding="utf-8")
        request.report(100, "FB2 pronto")
        return ConversionResult(output_path=out, converter=self.name, message="Convertido para FB2")


class EpubToDocxConverter(BaseConverter):
    """EPUB -> DOCX, keeping headings, paragraphs and images."""

    name = "epub_to_docx"
    priority = 60

    def can_handle(self, request: ConversionRequest) -> bool:
        return request.target_format == "docx" and request.detection.format == "epub"

    def supports_pair(self, source_fmt: str, target_fmt: str) -> bool:
        return source_fmt == "epub" and target_fmt == "docx"

    def run(self, request: ConversionRequest) -> ConversionResult:
        try:
            from docx import Document
            from docx.shared import Inches
        except ImportError as exc:  # pragma: no cover - declared dependency
            raise ConversionError("python-docx não está instalado.") from exc

        document = _load(request, with_images=True)
        doc = Document()
        for chapter_title, xhtml in epub_to_chapters(document, keep_markup=True):
            if chapter_title:
                doc.add_heading(chapter_title, level=1)
            _xhtml_into_docx(doc, xhtml, document, Inches)
        out = _output(request, "output.docx")
        doc.save(str(out))
        request.report(100, "DOCX pronto")
        return ConversionResult(output_path=out, converter=self.name, message="Convertido para DOCX")


# ---------------------------------------------------------------------------
# EPUB -> chapters (shared by TXT/FB2/DOCX)
# ---------------------------------------------------------------------------
def _load(request: ConversionRequest, *, with_images: bool = False) -> EpubDocument:
    try:
        return read_epub(request.source, include_images=with_images)
    except Exception as exc:  # noqa: BLE001 - one clear message for the user
        raise ConversionError(f"Não consegui ler o EPUB: {exc}") from exc


def epub_to_chapters(
    document: EpubDocument, *, keep_markup: bool = False
) -> list[tuple[str, str]]:
    """``[(chapter title, text or xhtml)]`` in spine order.

    Image ``src`` values are rewritten to their full path inside the EPUB, so
    callers can look the bytes up directly in ``document.items``.
    """
    chapters: list[tuple[str, str]] = []
    for item in document.documents():
        if not item.data:
            continue
        title, body = _document_to_content(
            item.data, base=item.name, document=document, keep_markup=keep_markup
        )
        if body.strip():
            chapters.append((title, body))
    return chapters or [(document.title or "Livro", "")]


def _document_to_content(
    raw: bytes, *, base: str, document: EpubDocument, keep_markup: bool
) -> tuple[str, str]:
    try:
        root = lxml_html.document_fromstring(raw)
    except etree.ParserError:
        return "", ""
    for element in root.iter():
        if isinstance(element.tag, str) and element.tag.lower() in {"script", "style"}:
            element.getparent().remove(element)
    _resolve_image_sources(root, base, document)
    heading = root.find(".//h1")
    title = (heading.text_content().strip() if heading is not None else "")[:200]
    if keep_markup:
        body_element = root.find(".//body")
        return title, _strip_outer(body_element if body_element is not None else root)
    text = root.text_content()
    return title, _clean_text(text)


def _resolve_image_sources(root, base: str, document: EpubDocument) -> None:
    """Point every ``img src`` at the real path inside the EPUB.

    Books are not always consistent: when the resolved path does not exist we
    fall back to the unique image with the same file name.
    """
    for image in root.iter("img"):
        src = image.get("src") or ""
        if not src or src.startswith(("http://", "https://", "data:")):
            continue
        resolved = document.resolve(base, src)
        if resolved not in document.items:
            basename = posixpath.basename(src)
            matches = [item for item in document.images() if posixpath.basename(item.name) == basename]
            if len(matches) == 1:
                resolved = matches[0].name
        image.set("src", resolved)


def _strip_outer(element) -> str:
    """Inner XHTML of an element (without its own tag)."""
    parts: list[str] = []
    if element.text:
        parts.append(escape(element.text))
    for child in element:
        parts.append(lxml_html.tostring(child, encoding="unicode"))
    return "".join(parts)


def _clean_text(text: str) -> str:
    lines = [line.strip() for line in text.splitlines()]
    return "\n\n".join(line for line in lines if line)


# ---------------------------------------------------------------------------
# Markup -> FB2 / DOCX
# ---------------------------------------------------------------------------
def _xhtml_to_fb2(xhtml: str) -> str:
    """XHTML fragment -> FB2 body markup."""
    try:
        fragment = lxml_html.fragment_fromstring(xhtml, create_parent="div")
    except etree.ParserError:
        return f"<p>{escape(_clean_text(xhtml))}</p>"
    parts: list[str] = []
    for element in fragment.iter():
        tag = element.tag if isinstance(element.tag, str) else ""
        text = (element.text or "").strip()
        if not text:
            continue
        if tag in {"h1", "h2", "h3", "h4"}:
            parts.append(f"<subtitle>{escape(text)}</subtitle>")
        elif tag in {"p", "div", "li", "blockquote"}:
            parts.append(f"<p>{_inline_fb2(element)}</p>")
    return "".join(parts) or "<p></p>"


def _inline_fb2(element) -> str:
    parts: list[str] = []
    if element.text:
        parts.append(escape(element.text))
    for child in element:
        tag = child.tag if isinstance(child.tag, str) else ""
        inner = _inline_fb2(child)
        if tag in {"em", "i"}:
            parts.append(f"<emphasis>{inner}</emphasis>")
        elif tag in {"strong", "b"}:
            parts.append(f"<strong>{inner}</strong>")
        elif tag == "br":
            parts.append("<empty-line/>")
        else:
            parts.append(inner)
    return "".join(parts)


def _xhtml_into_docx(doc, xhtml: str, document: EpubDocument, inches) -> None:
    try:
        fragment = lxml_html.fragment_fromstring(xhtml, create_parent="div")
    except etree.ParserError:
        doc.add_paragraph(_clean_text(xhtml))
        return
    for element in fragment.iter():
        tag = element.tag if isinstance(element.tag, str) else ""
        if tag in {"h1", "h2", "h3", "h4"}:
            text = (element.text_content() or "").strip()
            if text:
                doc.add_heading(text, level=min(3, int(tag[1])))
        elif tag == "img":
            _add_docx_image(doc, element, document, inches)
        elif tag in {"p", "div", "blockquote"}:
            text = (element.text_content() or "").strip()
            if text:
                doc.add_paragraph(text)
        elif tag == "li":
            text = (element.text_content() or "").strip()
            if text:
                doc.add_paragraph(text, style="List Bullet")


def _add_docx_image(doc, element, document: EpubDocument, inches) -> None:
    src = element.get("src") or ""
    if not src:
        return
    name = posixpath.normpath(src.lstrip("/"))
    item = document.items.get(name)
    if item is None or not item.data:
        return
    try:
        doc.add_picture(io.BytesIO(item.data), width=inches(4.5))
    except Exception:  # noqa: BLE001 - an image must not break the conversion
        return
