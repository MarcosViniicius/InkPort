"""Conversões nativas: os formatos que antes exigiam o Calibre.

EPUB -> KEPUB / TXT / FB2 / DOCX / AZW3 / MOBI e a volta para EPUB, tudo com
pacotes Python. O teste monta o próprio EPUB (não depende do acervo do usuário).

Run with:  python tests/native_formats.py
"""

from __future__ import annotations

import io
import os
import posixpath
import re
import shutil
import sys
import tempfile
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

WORKDIR = Path(tempfile.mkdtemp(prefix="opds_native_"))
os.environ["DATA_DIR"] = str(WORKDIR)
os.environ["SECRET_KEY"] = "test-secret-key"

PASSED: list[str] = []
FAILED: list[str] = []

#: Frase presente no livro de teste (com acento, para pegar erro de encoding).
PHRASE = "atenção"


def check(name: str, condition: bool, detail: str = "") -> None:
    if condition:
        PASSED.append(name)
        print(f"  ok   {name}")
    else:
        FAILED.append(f"{name}: {detail}")
        print(f"  FAIL {name} {detail}")


def _rewrite_entry(epub: Path, entry: str, transform) -> None:
    """Replace a single file inside a ZIP (used to build a broken fixture)."""
    with zipfile.ZipFile(epub) as source:
        nomes = source.namelist()
        dados = {nome: source.read(nome) for nome in nomes}
    dados[entry] = transform(dados[entry].decode("utf-8", "replace")).encode("utf-8")
    with zipfile.ZipFile(epub, "w", zipfile.ZIP_DEFLATED) as target:
        for nome in nomes:
            target.writestr(nome, dados[nome])


def _broken_refs(epub: Path) -> list[str]:
    """Local ``src``/``href`` values that do not resolve to an entry inside the EPUB.

    Readers resolve these relative to the file, so anything that lands outside
    the container shows up as a broken image/link.
    """
    problemas: list[str] = []
    with zipfile.ZipFile(epub) as zf:
        nomes = set(zf.namelist())
        for nome in nomes:
            if not nome.lower().endswith((".xhtml", ".html", ".htm")):
                continue
            texto = zf.read(nome).decode("utf-8", "replace")
            for valor in re.findall(r'(?:src|href)="([^"]+)"', texto):
                if not valor or valor.startswith(("http", "data:", "#", "mailto:")):
                    continue
                caminho = valor.split("#", 1)[0]
                if not caminho:
                    continue
                alvo = posixpath.normpath(
                    posixpath.join(posixpath.dirname(nome), caminho.lstrip("/"))
                )
                if alvo not in nomes:
                    problemas.append(f"{nome} -> {valor}")
    return problemas


def build_source_epub(path: Path) -> Path:
    """A small two-chapter EPUB with one image, written by our own builder."""
    from PIL import Image

    from app.converters.epub import EpubMeta
    from app.converters.epub.text_builder import write_text_epub

    buffer = io.BytesIO()
    Image.new("RGB", (60, 90), (200, 30, 30)).save(buffer, "PNG")
    chapters = [
        (
            "Introdução",
            f"<h1>Introdução</h1><p>A {PHRASE} é tudo o que você precisa.</p>"
            "<p>Segundo parágrafo, com ação e coração.</p>",
        ),
        (
            "Conclusão",
            "<h1>Conclusão</h1><p>Fim do livro de teste.</p>"
            '<p><img src="../images/figura.png" alt="figura"/></p>',
        ),
    ]
    write_text_epub(
        path,
        meta=EpubMeta(title="Livro de teste", author="Fulano de Tal", language="pt-BR", identifier="teste"),
        chapters=chapters,
        images=[("figura.png", buffer.getvalue(), "image/png")],
    )
    return path


def epub_text(path: Path) -> str:
    with zipfile.ZipFile(path) as zf:
        return "".join(
            zf.read(name).decode("utf-8", "replace")
            for name in zf.namelist()
            if name.endswith((".xhtml", ".html"))
        )


def main() -> int:
    from app.converters.base import ConversionRequest
    from app.converters.catalog import targets_for
    from app.converters.registry import converter_names, select_converter
    from app.converters.strategies.ebook_to_epub import (
        DocxToEpubConverter,
        Fb2ToEpubConverter,
        KindleToEpubConverter,
        TextToEpubConverter,
    )
    from app.converters.strategies.epub_to_ebook import (
        EpubToDocxConverter,
        EpubToFb2Converter,
        EpubToKepubConverter,
        EpubToTextConverter,
    )
    from app.devices.registry import get_profile
    from app.library.detect import detect
    from app.metadata.model import BookMetadata

    source = build_source_epub(WORKDIR / "livro.epub")
    out = WORKDIR / "out"
    out.mkdir(parents=True, exist_ok=True)
    profile = get_profile("eink_generic")

    def request(target: str, src: Path = source, workdir: Path | None = None) -> ConversionRequest:
        return ConversionRequest(
            source=src,
            detection=detect(src),
            metadata=BookMetadata(title="Livro de teste", author="Fulano de Tal"),
            profile=profile,
            target_format=target,
            workdir=workdir or out,
            options={},
        )

    print("[registro: os formatos do Calibre agora têm conversor nativo]")
    names = converter_names()
    for expected in (
        "epub_to_kindle",
        "epub_to_kepub",
        "epub_to_pdf",
        "epub_to_docx",
        "epub_to_fb2",
        "epub_to_text",
        "text_to_epub",
        "fb2_to_epub",
        "docx_to_epub",
        "kindle_to_epub",
    ):
        check(f"conversor {expected} registrado", expected in names)

    targets = [item["format"] for item in targets_for(detect(source), profile)]
    check(
        "EPUB oferece kepub/azw3/mobi/pdf/docx/fb2/txt",
        {"kepub", "azw3", "mobi", "pdf", "docx", "fb2", "txt"} <= set(targets),
        str(targets),
    )
    for target in ("azw3", "kepub", "docx", "fb2", "txt", "pdf"):
        chosen = select_converter(request(target))
        check(
            f"{target} resolvido sem Calibre",
            chosen is not None and chosen.name != "calibre",
            chosen.name if chosen else "nenhum",
        )

    print("\n[EPUB -> TXT -> EPUB]")
    text = EpubToTextConverter().run(request("txt"))
    body = text.output_path.read_text(encoding="utf-8", errors="replace")
    check("TXT com conteúdo", len(body) > 100, f"{len(body)} chars")
    check("TXT preserva acentos", PHRASE in body and "coração" in body)
    back = TextToEpubConverter().run(request("epub", text.output_path, out / "from_txt"))
    check("TXT -> EPUB", PHRASE in epub_text(back.output_path))

    print("\n[EPUB -> FB2 -> EPUB]")
    fb2 = EpubToFb2Converter().run(request("fb2"))
    payload = fb2.output_path.read_text(encoding="utf-8")
    check("FB2 é XML com FictionBook", payload.startswith("<?xml") and "<FictionBook" in payload)
    from lxml import etree

    try:
        root = etree.fromstring(payload.encode("utf-8"))
        sections = root.findall(".//{http://www.gribuser.ru/xml/fictionbook/2.0}section")
        check("FB2 válido com seções", len(sections) >= 2, str(len(sections)))
    except etree.XMLSyntaxError as exc:
        check("FB2 válido com seções", False, str(exc)[:70])
    back = Fb2ToEpubConverter().run(request("epub", fb2.output_path, out / "from_fb2"))
    check("FB2 -> EPUB", PHRASE in epub_text(back.output_path))

    print("\n[EPUB -> DOCX -> EPUB]")
    docx = EpubToDocxConverter().run(request("docx"))
    from docx import Document

    document = Document(str(docx.output_path))
    check("DOCX com parágrafos", len(document.paragraphs) >= 3, str(len(document.paragraphs)))
    check("DOCX com a imagem", len(document.inline_shapes) >= 1, str(len(document.inline_shapes)))
    back = DocxToEpubConverter().run(request("epub", docx.output_path, out / "from_docx"))
    check("DOCX -> EPUB", PHRASE in epub_text(back.output_path))

    print("\n[EPUB -> KEPUB]")
    kepub = EpubToKepubConverter().run(request("kepub"))
    with zipfile.ZipFile(kepub.output_path) as zf:
        entries = zf.namelist()
        check("KEPUB: mimetype é a primeira entrada", entries[0] == "mimetype", entries[0])
        check(
            "KEPUB: mimetype sem compressão",
            zf.getinfo("mimetype").compress_type == zipfile.ZIP_STORED,
        )
        markup = "".join(
            zf.read(name).decode("utf-8", "replace")
            for name in entries
            if name.endswith((".xhtml", ".html"))
        )
    check("KEPUB: spans do Kobo", 'class="koboSpan"' in markup)
    check("KEPUB: ids sequenciais", 'id="kobo.1.1"' in markup)

    print("\n[EPUB -> AZW3 / MOBI -> EPUB (escritor e leitor nativos)]")
    from app.converters.native.kindle import epub_to_kindle, kindle_available

    if not kindle_available():
        check("escritor Kindle disponível", False, "Pillow ausente?")
    else:
        for target in ("azw3", "mobi"):
            produced = WORKDIR / f"livro.{target}"
            try:
                epub_to_kindle(
                    source, produced, target=target, title="Livro de teste", author="Fulano"
                )
                check(f"{target.upper()} gerado", produced.stat().st_size > 1_500,
                      f"{produced.stat().st_size} bytes")
                back = KindleToEpubConverter().run(
                    request("epub", produced, out / f"from_{target}")
                )
                markup = epub_text(back.output_path)
                check(f"{target.upper()} -> EPUB preserva o texto", PHRASE in markup)
            except Exception as exc:  # noqa: BLE001
                check(f"{target.upper()} round-trip", False, f"{type(exc).__name__}: {str(exc)[:80]}")

    print("\n[PDF com texto -> EPUB (reflow nativo)]")
    try:
        import pymupdf

        pdf = WORKDIR / "texto.pdf"
        document = pymupdf.open()
        page = document.new_page()
        page.insert_text((72, 96), "Introducao", fontsize=20)
        page.insert_text((72, 130), f"A {PHRASE} e tudo o que voce precisa.", fontsize=11)
        page.insert_text((72, 150), "Segundo paragrafo do documento.", fontsize=11)
        document.save(str(pdf))
        document.close()
        chosen = select_converter(request("epub", pdf))
        check("PDF de texto escolhe o reflow nativo", chosen is not None and chosen.name == "pdf_to_epub",
              chosen.name if chosen else "nenhum")
        result = chosen.run(request("epub", pdf, out / "from_pdf"))
        produced_text = epub_text(result.output_path)
        check("PDF -> EPUB gerado", result.output_path.stat().st_size > 1_000)
        check(
            "PDF -> EPUB recupera o texto",
            "Introducao" in produced_text and "Segundo paragrafo" in produced_text,
            produced_text[:120],
        )
    except ImportError:
        print("  (PyMuPDF ausente: pulando)")

    print("\n[leitor: AZW3/MOBI sem Calibre]")
    from types import SimpleNamespace

    from app.reader.base import ReaderContext
    from app.reader.cache import ReaderCache
    from app.reader.native_convert import NativeConvertHandler

    azw3_source = WORKDIR / "leitor.azw3"
    epub_to_kindle(source, azw3_source, target="azw3", title="Livro de teste", author="Fulano")
    detected = detect(azw3_source)
    check(
        "AZW3 é detectado como azw3 (não como mobi)",
        detected.format == "azw3",
        detected.format,
    )

    handler = NativeConvertHandler()
    cache = ReaderCache("book-reader", azw3_source)
    context = ReaderContext(
        book=SimpleNamespace(id="book-reader", title="Livro de teste"),
        path=azw3_source,
        detection=detected,
        format="azw3",
        cache=cache,
    )
    before = handler.manifest(context)
    check("leitor: manifest inicial pede preparo", before.get("state") == "preparing", before.get("state"))
    handler.prepare(context)
    after = handler.manifest(context)
    chapters = after.get("chapters") or []
    check("leitor: fica pronto", after.get("state") == "ready", after.get("state"))
    check("leitor: delega ao EPUB", after.get("handler") == "epub", str(after.get("handler")))
    check("leitor: tem capítulos", len(chapters) >= 1, str(len(chapters)))
    found_phrase = False
    for chapter in chapters[:4]:
        body = handler.chapter(context, chapter["key"]).body.decode("utf-8", "replace")
        if PHRASE in body:
            found_phrase = True
            break
    check("leitor: texto com acentos", found_phrase, f"{len(chapters)} capítulos verificados")

    print("\n[capa embutida no EPUB gerado]")
    from PIL import Image

    from app.converters.epub import EpubMeta
    from app.converters.epub.text_builder import read_cover, write_text_epub

    cover_file = WORKDIR / "capa.png"
    Image.new("RGB", (200, 300), (180, 30, 40)).save(cover_file, "PNG")
    cover = read_cover(cover_file)
    check("read_cover lê a imagem", cover is not None and cover[1], str(cover and cover[0]))

    capped = WORKDIR / "com-capa.epub"
    write_text_epub(
        capped,
        meta=EpubMeta(title="Com capa", author="Fulano", language="pt", identifier="capa"),
        chapters=[("Um", "<h1>Um</h1><p>texto</p>")],
        images=[],
        cover=cover,
    )
    with zipfile.ZipFile(capped) as zf:
        names = zf.namelist()
        opf = zf.read("OEBPS/content.opf").decode("utf-8", "replace")
    check("capa guardada no zip", any(n.endswith("cover.png") for n in names), str(names[:6]))
    check("página de capa existe", "OEBPS/text/cover.xhtml" in names)
    check("OPF marca cover-image", 'properties="cover-image"' in opf)
    check("OPF traz meta name=cover (EPUB2)", 'name="cover"' in opf)
    check("guide aponta para a capa", 'reference type="cover"' in opf)
    check(
        "capa é a primeira página da espinha",
        opf.find('idref="coverpage"') < opf.find('idref="c0"'),
    )

    print("\n[EPUB de imagens: capa e referências internas válidas]")
    from app.converters.epub.builder import write_epub
    from app.library.repairs import _fix_epub_image_paths

    pagina_a = WORKDIR / "pagina-a.png"
    pagina_b = WORKDIR / "pagina-b.png"
    Image.new("RGB", (480, 800), (245, 245, 245)).save(pagina_a, "PNG")
    Image.new("RGB", (480, 800), (210, 220, 235)).save(pagina_b, "PNG")

    imagens_epub = WORKDIR / "imagens.epub"
    write_epub(
        imagens_epub,
        images=[pagina_a, pagina_b],
        meta=EpubMeta(
            title="Mangá de teste", language="pt", identifier="imagens", cover_image=cover_file
        ),
        include_title_page=True,
    )
    quebradas = _broken_refs(imagens_epub)
    check("EPUB de imagens sem referências quebradas", not quebradas, str(quebradas[:3]))
    with zipfile.ZipFile(imagens_epub) as zf:
        capa_gerada = zf.read("OEBPS/cover.xhtml").decode("utf-8", "replace")
        opf_gerado = zf.read("OEBPS/content.opf").decode("utf-8", "replace")
    check(
        "a capa aponta para a imagem dentro do container",
        'src="images/cover.png"' in capa_gerada and "../images/cover" not in capa_gerada,
        capa_gerada[:170],
    )
    check(
        "o OPF declara o tipo real da capa",
        'href="images/cover.png" media-type="image/png"' in opf_gerado,
    )

    # O mesmo reparo conserta arquivos antigos do acervo e EPUBs de terceiros.
    antigo = WORKDIR / "antigo.epub"
    shutil.copy2(imagens_epub, antigo)
    _rewrite_entry(
        antigo,
        "OEBPS/cover.xhtml",
        lambda texto: texto.replace('src="images/cover.png"', 'src="../images/cover.png"'),
    )
    check("fixture antiga tem a referência errada", bool(_broken_refs(antigo)))
    check("o reparo reescreve a referência", _fix_epub_image_paths(antigo))
    check(
        "depois do reparo não sobra referência quebrada",
        not _broken_refs(antigo),
        str(_broken_refs(antigo)[:3]),
    )
    check("o reparo é idempotente", not _fix_epub_image_paths(antigo))

    sem_capa = WORKDIR / "sem-capa.epub"
    write_text_epub(
        sem_capa,
        meta=EpubMeta(title="Sem capa", author="", language="pt", identifier="sem"),
        chapters=[("Um", "<h1>Um</h1><p>texto</p>")],
        images=[],
    )
    with zipfile.ZipFile(sem_capa) as zf:
        opf_sem = zf.read("OEBPS/content.opf").decode("utf-8", "replace")
    check("sem capa não inventa metadados", "cover-image" not in opf_sem and 'name="cover"' not in opf_sem)

    print("\n[imagens dos capítulos apontam para ../images/]")
    com_imagem = WORKDIR / "com-imagem.epub"
    write_text_epub(
        com_imagem,
        meta=EpubMeta(title="Com imagem", author="", language="pt", identifier="img"),
        chapters=[("Um", '<p>texto</p><p><img src="images/foto.jpg" alt=""/></p>')],
        images=[("foto.jpg", b"\xff\xd8\xff\xe0teste", "image/jpeg")],
    )
    with zipfile.ZipFile(com_imagem) as zf:
        capitulo = zf.read("OEBPS/text/chapter_0001.xhtml").decode("utf-8", "replace")
        guardada = "OEBPS/images/foto.jpg" in zf.namelist()
    referencias = re.findall(r'<img[^>]+src="([^"]+)"', capitulo)
    check("capítulo referencia ../images/", referencias == ["../images/foto.jpg"], str(referencias))
    check("a imagem está no pacote", guardada)

    print("\n[reparo de EPUBs antigos com o caminho errado]")
    from app.library.repairs import _fix_epub_image_paths

    quebrado = WORKDIR / "quebrado.epub"
    shutil.copyfile(com_imagem, quebrado)
    _rewrite_entry(
        quebrado,
        "OEBPS/text/chapter_0001.xhtml",
        lambda texto: texto.replace('src="../images/', 'src="images/'),
    )
    with zipfile.ZipFile(quebrado) as zf:
        antes = zf.read("OEBPS/text/chapter_0001.xhtml").decode("utf-8", "replace")
    check("o arquivo de teste começa quebrado", 'src="images/foto.jpg"' in antes, antes[:80])
    check("o reparo mexeu no arquivo", _fix_epub_image_paths(quebrado))
    with zipfile.ZipFile(quebrado) as zf:
        depois = zf.read("OEBPS/text/chapter_0001.xhtml").decode("utf-8", "replace")
    check("caminho relativo ficou correto", 'src="../images/foto.jpg"' in depois, depois[:80])
    check("reparo é idempotente", not _fix_epub_image_paths(quebrado))

    shutil.rmtree(WORKDIR, ignore_errors=True)
    print(f"\n{len(PASSED)} passed, {len(FAILED)} failed")
    for failure in FAILED:
        print(f"  - {failure}")
    return 1 if FAILED else 0


if __name__ == "__main__":
    sys.exit(main())
